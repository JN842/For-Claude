"""Daily snapshot to Excel, and point-in-time backtests.

The API returns each event's full history including its announce date, so
any past ``as_of`` can be rebuilt from today's download.  The dated Excel
file written by :func:`run_snapshot` is the record of what was forecast
and which membership bet was used on that day.
"""

from pathlib import Path

import numpy as np
import pandas as pd

from dividend_model import (
    MODEL_VERSION,
    ModelParams,
    assign_slots,
    build_event_table,
    prepare_events,
)
from set50 import (
    default_instruments,
    extend_membership,
    index_period,
    instrument_dividends,
    instrument_table,
    membership_changes,
)


def _ensure_prepared(events, params):
    if {"xd_date", "dps", "announce_date"}.issubset(events.columns):
        return events
    return prepare_events(events, params.imputed_announce_days)


def _horizon(instruments, as_of, holidays):
    table = instrument_table(instruments, as_of, holidays)
    return table["xd_until"].max()


def forecast_snapshot(events, as_of, membership, instruments=None,
                      holidays=(), params=None):
    """Event table and per-instrument dividends as known on ``as_of``."""
    params = params or ModelParams()
    as_of = pd.Timestamp(as_of).normalize()
    events = _ensure_prepared(events, params)
    instruments = list(instruments or default_instruments(as_of, holidays))
    horizon_end = _horizon(instruments, as_of, holidays)
    model = build_event_table(
        events,
        as_of=as_of,
        horizon_end=horizon_end,
        params=params,
        universe=membership.index,
    )
    carry = instrument_dividends(
        model["events"], membership, instruments, as_of, holidays
    )
    return {
        "as_of": as_of,
        "params": params,
        "instruments": carry["instruments"],
        "matrix": carry["matrix"],
        "detail": carry["detail"],
        "events": model["events"],
        "review": model["review"],
        "slots": model["slots"],
        "membership": extend_membership(
            membership, index_period(horizon_end)
        ),
    }


def _run_info(snapshot, scenario, holidays):
    rows = [
        ("as_of", snapshot["as_of"].date()),
        ("scenario", scenario),
        ("model_version", MODEL_VERSION),
        ("generated_at", pd.Timestamp.now().floor("s")),
        ("holidays_supplied", len(list(holidays))),
        ("symbols_in_universe", len(snapshot["membership"].index)),
        ("confirmed_events", int(snapshot["events"]["status"].eq("CONFIRMED").sum())),
        ("forecast_events", int(snapshot["events"]["status"].eq("FORECAST").sum())),
        ("review_rows", len(snapshot["review"])),
    ]
    rows += [(f"param.{k}", v) for k, v in snapshot["params"].as_dict().items()]
    return pd.DataFrame(rows, columns=["item", "value"])


def write_excel(snapshot, path, scenario="base", holidays=()):
    """Write the snapshot workbook.  Sheets, in order:

    Carry      symbols x instruments, expected THB/share (constituents only)
    Instruments  expiry dates and the XD range each instrument covers
    Detail     every event x instrument with probability, P(in range), member
    Events     confirmed + forecast cash events
    Membership the scenario used, plus a list of assumed changes
    Review     slots/symbols the model could not forecast
    Run_Info   as_of, scenario, parameters
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    membership = snapshot["membership"].astype(int)
    changes = membership_changes(snapshot["membership"])
    sheets = [
        ("Carry", snapshot["matrix"], True),
        ("Instruments", snapshot["instruments"], False),
        ("Detail", snapshot["detail"], False),
        ("Events", snapshot["events"], False),
        ("Membership", membership, True),
        ("Review", snapshot["review"], False),
        ("Run_Info", _run_info(snapshot, scenario, holidays), False),
    ]
    with pd.ExcelWriter(path, engine="openpyxl",
                        date_format="YYYY-MM-DD",
                        datetime_format="YYYY-MM-DD") as writer:
        for name, frame, keep_index in sheets:
            frame.to_excel(writer, sheet_name=name, index=keep_index)
        if not changes.empty:
            changes.to_excel(
                writer,
                sheet_name="Membership",
                startcol=len(membership.columns) + 3,
                index=False,
            )
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "B2"
            for column in sheet.columns:
                width = max(len(str(cell.value or "")) for cell in column)
                sheet.column_dimensions[column[0].column_letter].width = min(
                    max(10, width + 2), 40
                )
    return path


def read_holidays(path, sheet="Other", column="A", first_row=2):
    """SET holidays typed as dates in one column of a workbook.

    Reads the same list the trading workbook uses for its own expiry
    formulas, so both sides agree on last trading days.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        values = [
            row[0]
            for row in workbook[sheet].iter_rows(
                min_row=first_row,
                min_col=ord(column.upper()) - 64,
                max_col=ord(column.upper()) - 64,
                values_only=True,
            )
        ]
    finally:
        workbook.close()
    dates = pd.to_datetime(
        [value for value in values if value not in (None, "")],
        errors="coerce",
    )
    return sorted(set(dates.dropna().normalize()))


LINK_COLUMNS = {
    "carry": ["symbol", "instrument", "index_period", "expected_dps", "key"],
    "universe": ["universe"],
    "membership": ["member_symbol", "member_period", "is_member", "member_key"],
    "info": ["item", "value"],
}


def link_tables(snapshot, scenario="base"):
    """The blocks of the fixed-layout ``Link`` sheet for the trading workbook.

    carry: expected DPS per symbol, instrument and index half-year (the
    half-year of the XD date), so the workbook can apply the basket of
    that half-year.  Constituents only.
    """
    detail = snapshot["detail"]
    carry = (
        detail.loc[detail["is_member"]]
        .groupby(["symbol", "instrument", "index_period"], as_index=False)
        ["contribution"].sum()
        .rename(columns={"contribution": "expected_dps"})
    )
    membership = snapshot["membership"]
    member_long = (
        membership.astype(int)
        .rename_axis("member_symbol")
        .reset_index()
        .melt(id_vars="member_symbol", var_name="member_period",
              value_name="is_member")
    )
    # Text keys let the workbook use fast SUMIFS lookups.
    carry["key"] = (
        carry["symbol"] + "|" + carry["instrument"] + "|" + carry["index_period"]
    )
    member_long["member_key"] = (
        member_long["member_symbol"] + "|" + member_long["member_period"]
    )
    info = pd.DataFrame(
        [
            ("as_of", snapshot["as_of"]),
            ("scenario", scenario),
            ("model_version", MODEL_VERSION),
            ("generated_at", pd.Timestamp.now().floor("s")),
        ],
        columns=LINK_COLUMNS["info"],
    )
    return {
        "carry": carry[LINK_COLUMNS["carry"]],
        "universe": pd.DataFrame({"universe": list(membership.index)}),
        "membership": member_long[LINK_COLUMNS["membership"]],
        "info": info,
    }


def write_excel_link(snapshot, path, scenario="base"):
    """Write ``Link``: A:E carry, F universe, H:K membership, L:M info.

    The trading workbook reads these fixed columns (paste or Power Query
    into its ``Python`` sheet at A1), so the layout must not change.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tables = link_tables(snapshot, scenario)
    with pd.ExcelWriter(path, engine="openpyxl", date_format="YYYY-MM-DD",
                        datetime_format="YYYY-MM-DD") as writer:
        for name, start_col in [("carry", 0), ("universe", 5),
                                ("membership", 7), ("info", 11)]:
            tables[name].to_excel(writer, sheet_name="Link",
                                  startcol=start_col, index=False)
    return path


def run_snapshot(raw, as_of, membership, instruments=None, holidays=(),
                 params=None, output_dir="dividend_model_output",
                 scenario="base", link_name="dividend_latest.xlsx"):
    """Build today's snapshot and write two files to ``output_dir``:

    ``dividend_<as_of>_<scenario>.xlsx``  the dated record of this run
    ``link_name``  fixed-name ``Link`` sheet read by the trading workbook
    (overwritten each run; pass ``None`` to skip)
    """
    snapshot = forecast_snapshot(
        raw, as_of, membership, instruments, holidays, params
    )
    path = Path(output_dir) / (
        f"dividend_{snapshot['as_of']:%Y%m%d}_{scenario}.xlsx"
    )
    snapshot["path"] = write_excel(snapshot, path, scenario, holidays)
    if link_name:
        snapshot["link_path"] = write_excel_link(
            snapshot, Path(output_dir) / link_name, scenario
        )
    return snapshot


# ---------------------------------------------------------------------------
# Backtest
# ---------------------------------------------------------------------------

def _actual_events(events, as_of, horizon_end):
    actual = events.loc[
        events["xd_date"].gt(as_of) & events["xd_date"].le(horizon_end)
    ].copy()
    actual["slot"] = "ACTUAL"
    actual["status"] = "CONFIRMED"
    actual["probability"] = 1.0
    actual["expected_dps"] = actual["dps"]
    actual["xd_window_start"] = actual["xd_date"]
    actual["xd_window_end"] = actual["xd_date"]
    actual["event_key"] = (
        actual["symbol"] + "|" + actual["xd_date"].dt.strftime("%Y-%m-%d")
    )
    return actual


def _event_comparison(snapshot, events, horizon_end):
    """Match each forecast to the actual event in the same slot and year.

    Events already announced on ``as_of`` were inputs, not forecasts, and
    are left out.
    """
    as_of = snapshot["as_of"]
    forecasts = snapshot["events"].loc[
        snapshot["events"]["status"].eq("FORECAST")
    ].copy()
    forecasts["slot_year"] = (
        forecasts["event_key"].str.split("|").str[2].astype(int)
    )

    actual = events.loc[
        events["xd_date"].gt(as_of)
        & events["xd_date"].le(horizon_end)
        & events["announce_date"].gt(as_of)
    ]
    actual = assign_slots(actual, snapshot["slots"], snapshot["params"])
    actual["slot"] = actual["slot"].fillna("UNSLOTTED")
    actual["slot_year"] = actual["slot_year"].fillna(0).astype(int)
    actual = actual.groupby(
        ["symbol", "slot", "slot_year"], as_index=False
    ).agg(actual_xd=("xd_date", "min"), actual_dps=("dps", "sum"))

    merged = forecasts.merge(
        actual, on=["symbol", "slot", "slot_year"], how="outer"
    )
    merged["outcome"] = np.select(
        [
            merged["xd_date"].notna() & merged["actual_xd"].notna(),
            merged["xd_date"].notna(),
        ],
        ["HIT", "NO_EVENT"],
        default="UNEXPECTED",
    )
    merged["xd_error_days"] = (merged["actual_xd"] - merged["xd_date"]).dt.days
    merged["dps_error"] = merged["dps"] - merged["actual_dps"]
    merged["in_window"] = merged["actual_xd"].between(
        merged["xd_window_start"], merged["xd_window_end"]
    )
    columns = [
        "symbol", "slot", "outcome", "probability", "xd_date", "actual_xd",
        "xd_error_days", "in_window", "dps", "actual_dps", "dps_error",
    ]
    return merged[columns].sort_values(
        ["symbol", "xd_date", "actual_xd"]
    ).reset_index(drop=True)


def backtest_snapshot(raw, as_of, membership, instruments=None, holidays=(),
                      params=None, actual_membership=None):
    """Forecast as of a past date and compare with what was actually paid.

    ``membership`` is the constituent bet used for the forecast; pass the
    real lists as ``actual_membership`` to also measure the membership
    error (it defaults to the bet, which isolates the dividend model).
    Only data announced on or before ``as_of`` reaches the forecast.

    Note: ``adjustedDPS`` is restated for later corporate actions, so DPS
    history carries a small look-ahead that the API does not let us undo.
    """
    params = params or ModelParams()
    events = _ensure_prepared(raw, params)
    snapshot = forecast_snapshot(
        events, as_of, membership, instruments, holidays, params
    )
    as_of = snapshot["as_of"]
    instruments = list(snapshot["instruments"]["instrument"])
    horizon_end = snapshot["instruments"]["xd_until"].max()
    if horizon_end > events["xd_date"].max():
        raise ValueError(
            f"Data ends {events['xd_date'].max():%Y-%m-%d}; choose an as_of "
            f"whose instruments expire by then (horizon {horizon_end:%Y-%m-%d})"
        )

    actual_membership = (
        membership if actual_membership is None else actual_membership
    )
    realised = instrument_dividends(
        _actual_events(events, as_of, horizon_end),
        actual_membership,
        instruments,
        as_of,
        holidays,
    )

    forecast_long = snapshot["matrix"].stack().rename("forecast")
    actual_long = realised["matrix"].stack().rename("actual")
    carry = pd.concat([forecast_long, actual_long], axis=1).fillna(0.0)
    carry.index.names = ["symbol", "instrument"]
    carry = carry.reset_index()
    carry["error"] = carry["forecast"] - carry["actual"]
    carry.insert(0, "as_of", as_of)

    return {
        "as_of": as_of,
        "carry": carry,
        "events": _event_comparison(snapshot, events, horizon_end).assign(
            as_of=as_of
        ),
        "snapshot": snapshot,
    }


def backtest_many(raw, as_of_dates, membership, instruments_for=None,
                  holidays=(), params=None, actual_membership=None):
    """Run :func:`backtest_snapshot` for several dates and stack the results.

    ``instruments_for(as_of)`` chooses the instruments for each date
    (default: next two quarterly contracts and their spread).
    """
    params = params or ModelParams()
    events = _ensure_prepared(raw, params)
    instruments_for = instruments_for or (
        lambda as_of: default_instruments(as_of, holidays)
    )
    carry_frames, event_frames = [], []
    for as_of in as_of_dates:
        result = backtest_snapshot(
            events,
            as_of,
            membership,
            instruments_for(pd.Timestamp(as_of)),
            holidays,
            params,
            actual_membership,
        )
        carry_frames.append(result["carry"])
        event_frames.append(result["events"])
    carry = pd.concat(carry_frames, ignore_index=True)
    event_results = pd.concat(event_frames, ignore_index=True)
    summary = (
        carry.groupby(["as_of", "instrument"], as_index=False)
        .agg(
            forecast_dps_sum=("forecast", "sum"),
            actual_dps_sum=("actual", "sum"),
            abs_error_sum=("error", lambda error: error.abs().sum()),
        )
    )
    return {"carry": carry, "events": event_results, "summary": summary}
