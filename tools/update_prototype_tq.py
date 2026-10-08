"""Wire the Python dividend output into the trading workbook (prototype_TQ.xlsx).

    python tools/update_prototype_tq.py prototype_TQ.xlsx prototype_TQ_linked.xlsx

Adds/rebuilds, leaving the other sheets untouched:

Python       paste target for dividend_latest.xlsx > Link (A1); left empty
Basket Proj  projected baskets for the next 4 quarters, each with the
             membership of its index half-year, same method as Q Rebalance
             (unrounded shares)
Dividend     expected DPS per symbol x Monitor series x XD quarter, times the
             basket of that quarter, / (multiplier x lots) = index points
Monitor      Div pts, fair value at risk-free and borrowing rates, and a
             calendar-spread table; holiday ranges in M7:M10 made absolute

openpyxl drops Excel-2010 conditional-format extensions, so the Monitor
sheet's <extLst> is copied back from the source file after saving.
"""

import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FIRST, LAST = 9, 88          # symbol rows in Dividend and Basket Proj
SERIES_ROWS = [7, 8, 9, 10]  # Monitor rows holding the rolling series

BASE = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", size=11, bold=True)
INPUT = Font(name="Calibri", size=11, color="FFFF0000")  # workbook convention
HEAD_FILL = PatternFill("solid", fgColor="FFD9E1F2")
THIN = Side(style="thin", color="FFBFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
DPS = '0.0000;-0.0000;"-"'
PTS = '0.00'


def _header(cell, text):
    cell.value = text
    cell.font = BOLD
    cell.fill = HEAD_FILL
    cell.border = BOX
    cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _put(ws, ref, value, fmt=None, font=BASE):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fmt:
        cell.number_format = fmt
    return cell


def _match_text(range_ref, key):
    # Compare as text: Excel turns the symbol TRUE into a boolean.
    return f'--({range_ref}&""={key}&"")'


def build_python(ws):
    ws.delete_rows(1, ws.max_row)
    ws["A1"].comment = Comment(
        "Paste dividend_latest.xlsx > Link here at A1 (values), or load it "
        "with Power Query to A1. Columns: A:E carry, F universe, "
        "H:K membership, L:M info. Do not type in this sheet.",
        "python",
    )


QUARTERS = 5                 # current quarter + 4 projected
BLOCK = 6                    # columns per projected-quarter block (5 + gap)


def _col(index):
    return get_column_letter(index)


def _proj_block(k):
    """Column letters of projected quarter k (1..4) in Basket Proj."""
    start = 8 + (k - 1) * BLOCK  # H, N, T, Z
    return [_col(start + i) for i in range(5)]


def _shares_match(sheet, key_col, value_col, r, first=FIRST, last=LAST):
    return (f"=IF($B{r}=\"\",0,SUMPRODUCT("
            f"{_match_text(f'{chr(39)}{sheet}{chr(39)}!${key_col}${first}:${key_col}${last}', '$B' + str(r))},"
            f"'{sheet}'!${value_col}${first}:${value_col}${last}))")


def build_basket_proj(wb):
    for old in ("Basket Next", "Basket Proj"):
        if old in wb.sheetnames:
            del wb[old]
    ws = wb.create_sheet("Basket Proj", wb.sheetnames.index("Q Rebalance") + 1)
    ws["B1"] = ("Projected baskets for the next 4 quarters "
                "(Q Rebalance method; members from Python per index half-year)")
    ws["B1"].font = BOLD

    inputs = [
        ("B3", "Price date", "C3", "='Q Rebalance'!$K$3", "dd/mm/yyyy", True),
        ("B4", "SET50 close", "C4",
         '=RTD("profrtd",,"history","SET50,DAILY,PCLOSE,E="&TEXT($C$3,"ddmmyyyy")'
         '&TEXT("","hhmm")&"",1)', "#,##0.00", False),
        ("B5", "Weight cap", "C5", 0.1, "0.0%", True),
        ("B6", "Basket value (THB)", "C6", "=C4*Dividend!$C$3*Dividend!$C$4", NUM, False),
    ]
    for label_ref, label, value_ref, value, fmt, is_input in inputs:
        _put(ws, label_ref, label)
        _put(ws, value_ref, value, fmt, INPUT if is_input else BASE)
    ws["C3"].comment = Comment(
        "Closing prices for all projected weights. Defaults to the Q "
        "Rebalance date; type a date to override.", "python")
    ws["C5"].comment = Comment(
        "Same 10% cap Q Rebalance applies to DELTA, applied to any member "
        "above the cap in one pass.", "python")

    for col, text in zip("BCDEF", ["Symbol", "Close", "Shares (Q Rebalance)",
                                   "Shares override", "Shares used"]):
        _header(ws[f"{col}8"], text)
    ws["E8"].comment = Comment(
        "Type listed shares for a stock that is not in Q Rebalance "
        "(e.g. a stock you bet will enter). Used in every quarter.", "python")

    for k in range(1, QUARTERS):
        m, mc, w, cap, q = _proj_block(k)
        _put(ws, f"{m}3", "Quarter", font=BOLD)
        _put(ws, f"{mc}3", f"=Dividend!{_col(3 + k)}$8", font=BOLD)
        _put(ws, f"{m}4", "Index half-year")
        _put(ws, f"{mc}4",
             f'=LEFT({mc}3,4)&"H"&IF(VALUE(RIGHT({mc}3,1))<=2,1,2)')
        _put(ws, f"{m}5", "Members")
        _put(ws, f"{mc}5", f"=SUM({m}{FIRST}:{m}{LAST})", "0")
        _put(ws, f"{m}6", "Missing shares/price")
        _put(ws, f"{mc}6",
             f"=COUNTIFS({m}{FIRST}:{m}{LAST},1,$F${FIRST}:$F${LAST},0)"
             f"+COUNTIFS({m}{FIRST}:{m}{LAST},1,$C${FIRST}:$C${LAST},0)", "0",
             INPUT)
        _put(ws, f"{mc}7", f"=SUM({mc}{FIRST}:{mc}{LAST})", NUM)
        _put(ws, f"{cap}7", f"=SUM({cap}{FIRST}:{cap}{LAST})", NUM)
        for col, text in zip((m, mc, w, cap, q), ["Member", "Mkt Cap (M)",
                                                  "Raw weight", "Capped (M)",
                                                  "Q Stock"]):
            _header(ws[f"{col}8"], text)
        for col, width in zip((m, mc, w, cap, q), [9, 13, 10, 13, 12]):
            ws.column_dimensions[col].width = width

    for row in range(FIRST, LAST + 1):
        link_row = row - FIRST + 2
        r = str(row)
        _put(ws, "B" + r, f'=IF(Python!$F${link_row}="","",Python!$F${link_row}&"")')
        _put(ws, "C" + r,
             f'=IF($B{r}="",0,IFERROR(RTD("profrtd",,"history",$B{r}&",DAILY,PCLOSE,E="'
             f'&TEXT($C$3,"ddmmyyyy")&TEXT("","hhmm")&"",1),0))', "#,##0.00")
        _put(ws, "D" + r, _shares_match("Q Rebalance", "J", "L", row, 8, 57), NUM)
        ws["E" + r].font = INPUT
        ws["E" + r].number_format = NUM
        _put(ws, "F" + r, f'=IF(E{r}<>"",E{r},D{r})', NUM)
        for k in range(1, QUARTERS):
            m, mc, w, cap, q = _proj_block(k)
            _put(ws, m + r,
                 f'=IF($B{r}="",0,SUMIFS(Python!$J:$J,Python!$K:$K,$B{r}&"|"&${mc}$4))', "0")
            _put(ws, mc + r, f"=IF({m}{r}=1,$C{r}*$F{r}/1000000,0)", NUM)
            _put(ws, w + r, f"=IF(${mc}$7=0,0,{mc}{r}/${mc}$7)", "0.00%")
            _put(ws, cap + r,
                 f"=IF({w}{r}>$C$5,$C$5/(1-COUNTIF(${w}${FIRST}:${w}${LAST},\">\"&$C$5)*$C$5)"
                 f"*SUMIF(${w}${FIRST}:${w}${LAST},\"<=\"&$C$5,${mc}${FIRST}:${mc}${LAST}),{mc}{r})",
                 NUM)
            _put(ws, q + r,
                 f"=IF(AND($C{r}>0,${cap}$7>0),{cap}{r}/${cap}$7*$C$6/$C{r},0)", "#,##0.00")

    for col, width in zip("ABCDEFG", [2, 12, 9, 18, 16, 16, 2]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "C9"


def build_dividend(wb):
    ws = wb["Dividend"]
    ws.delete_rows(1, ws.max_row)
    _put(ws, "B1", "Expected dividend points per series "
                   "(Python carry x basket of the XD quarter)", font=BOLD)

    _put(ws, "B3", "Multiplier")
    _put(ws, "C3", 200, "0", INPUT)
    _put(ws, "B4", "Lots in basket")
    _put(ws, "C4", 50, "0", INPUT)
    _put(ws, "B5", "Python as of")
    _put(ws, "C5", '=IF(Python!$M$2="","no data",Python!$M$2)', "dd/mm/yyyy")
    _put(ws, "D5",
         '=IF(Python!$M$2="","Paste Link into Python!A1",'
         'IF(INT(Python!$M$2)=TODAY(),"","STALE: rerun Python"))', font=INPUT)
    _put(ws, "B6", "Q Rebalance date")
    _put(ws, "C6", "='Q Rebalance'!$K$3", "dd/mm/yyyy")
    _put(ws, "D6", '=IF(TODAY()-C6>100,"Q Rebalance older than a quarter","")',
         font=INPUT)

    for ref, text in [("F2", "Series"), ("F3", "Div points"),
                      ("F4", "Current quarter"), ("F5", "Later quarters")]:
        _header(ws[ref], text)

    shares_cols = [_col(3 + j) for j in range(QUARTERS)]          # C..G
    _header(ws["B7"], "Series ->")
    _header(ws["B8"], "Symbol")
    for j, col in enumerate(shares_cols):
        _header(ws[f"{col}7"], "Shares")
        if j == 0:
            formula = '=YEAR(TODAY())&"Q"&ROUNDUP(MONTH(TODAY())/3,0)'
        else:
            prev = shares_cols[j - 1]
            formula = (f'=IF(RIGHT({prev}8,1)="4",(LEFT({prev}8,4)+1)&"Q1",'
                       f'LEFT({prev}8,5)&(RIGHT({prev}8,1)+1))')
        _put(ws, f"{col}8", formula, font=BOLD)
    ws["C7"].comment = Comment("Current quarter: Q Rebalance column P (unrounded).", "python")
    ws["D7"].comment = Comment("Later quarters: Basket Proj Q Stock of that quarter.", "python")

    def detail_col(k, j):
        return _col(9 + QUARTERS * k + j)                          # I..AB

    lots = "($C$3*$C$4)"
    for k, series_row in enumerate(SERIES_ROWS):
        summary_col = "GHIJ"[k]
        _put(ws, f"{summary_col}2", f"=Monitor!$I${series_row}", font=BOLD)
        products = [
            f"SUMPRODUCT({detail_col(k, j)}${FIRST}:{detail_col(k, j)}${LAST},"
            f"${shares_cols[j]}${FIRST}:${shares_cols[j]}${LAST})"
            for j in range(QUARTERS)
        ]
        _put(ws, f"{summary_col}3", f"=({'+'.join(products)})/{lots}", PTS, BOLD)
        _put(ws, f"{summary_col}4", f"={products[0]}/{lots}", PTS)
        _put(ws, f"{summary_col}5", f"=({'+'.join(products[1:])})/{lots}", PTS)
        for j in range(QUARTERS):
            c = detail_col(k, j)
            _put(ws, f"{c}7", f"=Monitor!$I${series_row}", font=BOLD)
            _put(ws, f"{c}8", f"=${shares_cols[j]}$8", font=BOLD)
            ws[f"{c}7"].fill = ws[f"{c}8"].fill = HEAD_FILL
            ws.column_dimensions[c].width = 10

    check = _col(9 + QUARTERS * len(SERIES_ROWS) + 1)               # AD
    _header(ws[f"{check}8"], "Check")
    proj_q = [_proj_block(k)[4] for k in range(1, QUARTERS)]
    for row in range(FIRST, LAST + 1):
        link_row = row - FIRST + 2
        r = str(row)
        _put(ws, "B" + r, f'=IF(Python!$F${link_row}="","",Python!$F${link_row}&"")')
        _put(ws, "C" + r, _shares_match("Q Rebalance", "J", "P", row, 8, 57), "#,##0")
        for j in range(1, QUARTERS):
            _put(ws, shares_cols[j] + r,
                 _shares_match("Basket Proj", "B", proj_q[j - 1], row), "#,##0")
        for k in range(len(SERIES_ROWS)):
            for j in range(QUARTERS):
                c = detail_col(k, j)
                _put(ws, c + r,
                     f'=IF($B{r}="",0,SUMIFS(Python!$D:$D,Python!$E:$E,'
                     f'$B{r}&"|"&{c}$7&"|"&{c}$8))', DPS)
        conditions = [
            "AND(" + "+".join(f"{detail_col(k, j)}{r}" for k in range(len(SERIES_ROWS)))
            + f">0,{shares_cols[j]}{r}=0)"
            for j in range(QUARTERS)
        ]
        _put(ws, check + r, f'=IF(OR({",".join(conditions)}),"DPS but no shares","")',
             font=INPUT)

    for col, width in zip("ABCDEFGH", [2, 12, 11, 11, 11, 11, 18, 2]):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions[check].width = 18
    ws.freeze_panes = "C9"


def build_monitor(ws):
    # Holiday list must not shift row by row.
    for row in SERIES_ROWS:
        cell = ws[f"M{row}"]
        cell.value = re.sub(r"Other!\$?A\$?\d+:\$?A\$?\d+", "Other!$A$2:$A$30", cell.value)

    for col, text in zip("PQR", ["Div pts", "Fair (Rf)", "Fair (Borrow)"]):
        _put(ws, f"{col}6", text)
    for k, row in enumerate(SERIES_ROWS):
        _put(ws, f"P{row}", f"=Dividend!{'GHIJ'[k]}$3", PTS)
        _put(ws, f"Q{row}", f"=$D$6*(1+$I$3*O{row})-P{row}", PTS)
        _put(ws, f"R{row}", f"=$D$6*(1+$J$3*O{row})-P{row}", PTS)
    _put(ws, "P5", '=Dividend!$D$5', font=INPUT)

    _put(ws, "I12", "Calendar spread (far - near)")
    for col, text in zip("IJKLMNO", ["Spread", "Bid", "Ask", "Mid", "Div pts",
                                      "Fair (Rf)", "Fair (Borrow)"]):
        _put(ws, f"{col}13", text)
    for i, (near, far) in enumerate(zip(SERIES_ROWS, SERIES_ROWS[1:])):
        row = 14 + i
        _put(ws, f"I{row}", f'=I{near}&MID(I{far},4,3)')
        _put(ws, f"J{row}", f"=J{far}-K{near}", PTS)   # sell far, buy near
        _put(ws, f"K{row}", f"=K{far}-J{near}", PTS)   # buy far, sell near
        _put(ws, f"L{row}", f"=L{far}-L{near}", PTS)
        _put(ws, f"M{row}", f"=P{far}-P{near}", PTS)
        _put(ws, f"N{row}", f"=Q{far}-Q{near}", PTS)
        _put(ws, f"O{row}", f"=R{far}-R{near}", PTS)
    ws["P6"].comment = Comment(
        "Expected dividend index points with XD after the Python as-of date "
        "and on or before this series' last trading day (Dividend sheet).",
        "python")
    ws["Q6"].comment = Comment("Spot x (1 + risk-free x t) - Div pts", "python")
    ws["R6"].comment = Comment("Spot x (1 + borrowing x t) - Div pts", "python")


def build_other(ws):
    if ws["A1"].value in (None, ""):
        ws["A1"] = "SET holidays (dates, A2:A30)"
        ws["A1"].font = BOLD


def _sheet_xml(path, sheet_name):
    """Zip path of a sheet's XML, whatever attribute order the writer used."""
    with zipfile.ZipFile(path) as z:
        workbook = z.read("xl/workbook.xml").decode()
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
    for tag in re.findall(r"<sheet\b[^>]*>", workbook):
        if re.search(rf'\bname="{re.escape(sheet_name)}"', tag):
            rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
            break
    else:
        raise KeyError(sheet_name)
    for tag in re.findall(r"<Relationship\b[^>]*>", rels):
        if re.search(rf'\bId="{rid}"', tag):
            target = re.search(r'\bTarget="([^"]+)"', tag).group(1)
            return "xl/" + target.lstrip("/").removeprefix("xl/")
    raise KeyError(rid)


def main(source, target):
    source, target = Path(source), Path(target)
    monitor_xml_src = _sheet_xml(source, "Monitor")
    wb = load_workbook(source)
    build_other(wb["Other"])
    build_python(wb["Python"])
    build_basket_proj(wb)
    build_dividend(wb)
    build_monitor(wb["Monitor"])
    wb.calculation.fullCalcOnLoad = True
    wb.save(target)
    with zipfile.ZipFile(source) as z:
        original = z.read(monitor_xml_src).decode()
    _restore_ext_lst_text(original, target, _sheet_xml(target, "Monitor"))


def _restore_ext_lst_text(original_xml, target, sheet_xml):
    match = re.search(r"<extLst>.*</extLst>", original_xml, re.S)
    if not match:
        return
    tmp = Path(tempfile.mkdtemp()) / "out.xlsx"
    with zipfile.ZipFile(target) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_xml:
                text = data.decode("utf-8")
                if "<extLst>" not in text:
                    text = text.replace("</worksheet>", match.group(0) + "</worksheet>")
                data = text.encode("utf-8")
            zout.writestr(item, data)
    shutil.move(tmp, target)


if __name__ == "__main__":
    main(*sys.argv[1:3])
